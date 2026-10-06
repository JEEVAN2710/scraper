"""Unit and integration tests for Excel export engine and FastAPI endpoints."""

import io
from unittest.mock import MagicMock, patch
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.api.routes import DEMO_COMPANIES
from app.export.excel_exporter import (
    generate_company_csv,
    generate_company_excel,
    generate_sector_master_excel,
)
from app.main import app


@pytest.fixture
def client():
    """FastAPI TestClient fixture."""
    return TestClient(app)


def test_excel_export_generation():
    """Verify Excel exporter generates single clean KPI Breakdown sheet by default, and auxiliary sheets when requested."""
    sample_company = DEMO_COMPANIES[0]  # TCS

    # 1. Default: Clean single-sheet export (as requested by user)
    clean_buf = generate_company_excel(sample_company)
    assert isinstance(clean_buf, io.BytesIO)
    clean_buf.seek(0)
    wb_clean = openpyxl.load_workbook(clean_buf)
    assert wb_clean.sheetnames == ["KPI Breakdown"]
    ws_kpi = wb_clean["KPI Breakdown"]
    assert "Tata Consultancy Services" in str(ws_kpi["A1"].value)

    # 2. Auxiliary sheets when explicitly requested
    excel_buffer = generate_company_excel(sample_company, include_auxiliary_sheets=True)
    assert isinstance(excel_buffer, io.BytesIO)
    excel_buffer.seek(0)

    wb = openpyxl.load_workbook(excel_buffer)
    sheet_names = wb.sheetnames

    assert "KPI Breakdown" in sheet_names
    assert "Company Overview" in sheet_names
    assert "Financial Metrics" in sheet_names
    assert "Risk Factors" in sheet_names
    assert "Document Registry" in sheet_names

    # Check Overview sheet content
    ws_ov = wb["Company Overview"]
    assert "Tata Consultancy Services" in str(ws_ov["A1"].value)

    # Check Financial Metrics sheet headers
    ws_fin = wb["Financial Metrics"]
    headers = [cell.value for cell in ws_fin[1]]
    assert "Period" in headers
    assert "Revenue" in headers
    assert "Net Profit" in headers
    assert "Operating Margin" in headers

    # Verify rows exist
    assert ws_fin.max_row >= 4


def test_sector_master_excel_generation():
    """Verify Sector Master Excel exporter generates MASTER tab, breakdown, and company tabs."""
    excel_buffer = generate_sector_master_excel(DEMO_COMPANIES)
    assert isinstance(excel_buffer, io.BytesIO)
    excel_buffer.seek(0)

    wb = openpyxl.load_workbook(excel_buffer)
    sheet_names = wb.sheetnames

    assert "MASTER" in sheet_names
    assert "Company Breakdown" in sheet_names
    # Check MASTER sheet title banner
    ws_master = wb["MASTER"]
    assert "MASTER" in str(ws_master["A1"].value)
    assert "Coworking Sector Review" in str(ws_master["A1"].value)

    # Check that sections are present
    all_col_a = [str(ws_master.cell(row=r, column=1).value) for r in range(1, ws_master.max_row + 1)]
    assert any("Revenue" in val for val in all_col_a)
    assert any("EBITDA" in val for val in all_col_a)
    assert any("Total" in val for val in all_col_a)

    # Check Company Breakdown sheet
    ws_breakdown = wb["Company Breakdown"]
    assert ws_breakdown.max_row >= 5


def test_csv_export_generation():
    """Verify CSV string generation with valid header and values."""
    sample_company = DEMO_COMPANIES[0]
    csv_str = generate_company_csv(sample_company)

    assert "period,revenue,revenue_growth,net_profit" in csv_str
    assert "FY2024" in csv_str


def test_api_health_endpoint(client):
    """Test GET /api/health endpoint."""
    res = client.get("/api/health")
    assert res.status_code == 200
    data = res.json()
    assert "status" in data
    assert "mysql" in data
    assert "ollama" in data


def test_api_companies_endpoint(client):
    """Test GET /api/companies endpoint."""
    res = client.get("/api/companies")
    assert res.status_code == 200
    companies = res.json()
    assert isinstance(companies, list)


def test_api_company_detail_endpoint(client):
    """Test GET /api/companies/{id} endpoint."""
    with patch("app.api.routes._get_company_from_db", return_value=DEMO_COMPANIES[0]):
        res = client.get("/api/companies/1")
        assert res.status_code == 200
        detail = res.json()
        assert detail["id"] == 1
        assert len(detail["financial_data"]) > 0
        assert len(detail["risks"]) > 0
        assert len(detail["documents"]) > 0


def test_api_excel_download_endpoint(client):
    """Test GET /api/export/excel/{id} download response."""
    with patch("app.api.routes._get_company_from_db", return_value=DEMO_COMPANIES[0]):
        res = client.get("/api/export/excel/1")
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "attachment; filename=" in res.headers["content-disposition"]
        assert len(res.content) > 1000


def test_api_sector_excel_download_endpoint(client):
    """Test GET /api/export/sector-excel download response."""
    with patch("app.api.routes._get_company_from_db", return_value=DEMO_COMPANIES[0]), \
         patch("app.api.routes.db_manager.get_cursor") as mock_cur:
        mock_cursor = MagicMock()
        mock_cursor.fetchall.return_value = [{"id": 1}]
        mock_cur.return_value.__enter__.return_value = (mock_cursor, MagicMock())
        res = client.get("/api/export/sector-excel")
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "attachment; filename=" in res.headers["content-disposition"]
        assert "Sector_Review_Master_" in res.headers["content-disposition"]
        assert len(res.content) > 1000


def test_api_csv_download_endpoint(client):
    """Test GET /api/export/csv/{id} download response."""
    with patch("app.api.routes._get_company_from_db", return_value=DEMO_COMPANIES[0]):
        res = client.get("/api/export/csv/1")
        assert res.status_code == 200
        assert "text/csv" in res.headers["content-type"]
        assert len(res.text) > 50


def test_api_pdf_download_endpoint(client):
    """Test GET /api/documents/{id}/download response."""
    with patch("app.api.routes.db_manager.get_cursor") as mock_cur:
        mock_cursor = MagicMock()
        mock_cursor.fetchone.return_value = {
            "id": 101,
            "company_id": 1,
            "file_name": "TCS_Annual_Report_FY24.pdf",
            "report_period": "FY2024",
            "local_path": "data/downloads/TCS_Annual_Report_FY24.pdf",
        }
        mock_cur.return_value.__enter__.return_value = (mock_cursor, MagicMock())
        res = client.get("/api/documents/101/download")
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/pdf"
        assert res.content.startswith(b"%PDF-1.4")


def test_api_ask_graph_rag_endpoint(client):
    """Test POST /api/ask endpoint with Graph RAG."""
    from app.config.settings import get_settings
    from app.pipeline.graphify import get_knowledge_graph
    kg = get_knowledge_graph(get_settings())
    kg.add_node("comp_1", "Tata Consultancy Services Ltd. (TCS)", "company", {"company_id": 1, "name": "Tata Consultancy Services Ltd.", "ticker": "TCS"})
    kg.add_node("fin_1", "FY2024: $29.1B", "metric", {"company_id": 1, "period": "FY2024", "revenue": 29080000000.0, "revenue_formatted": "$29.1B"})
    kg.add_edge("comp_1", "fin_1", "REPORTED_FINANCIALS")

    res = client.post(
        "/api/ask",
        json={"question": "What was TCS revenue in FY2024?", "company_id": 1},
    )
    assert res.status_code == 200
    data = res.json()
    assert "answer" in data
    assert "graph_nodes" in data
    assert "citations" in data
    assert "TCS" in str(data["graph_nodes"]) or "Tata" in str(data["graph_nodes"])
    assert len(data["answer"]) > 20


def test_api_ask_cross_company_comparison(client):
    """Test POST /api/ask cross-company comparison inquiry."""
    from app.config.settings import get_settings
    from app.pipeline.graphify import get_knowledge_graph
    kg = get_knowledge_graph(get_settings())
    kg.add_node("comp_apple", "Apple Inc. (AAPL)", "company", {"name": "Apple Inc.", "ticker": "AAPL"})
    kg.add_node("comp_msft", "Microsoft Corporation (MSFT)", "company", {"name": "Microsoft Corporation", "ticker": "MSFT"})

    res = client.post(
        "/api/ask",
        json={"question": "Compare Apple and Microsoft profit margins"},
    )
    assert res.status_code == 200
    data = res.json()
    assert len(data["graph_nodes"]) >= 2
    assert "answer" in data
